from pathlib import Path
import shutil
import json
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

out = Path(r'C:\dev\XuChe\幼儿受伤处理记录\08-律师咨询')
qa = Path(r'C:\dev\XuChe\tmp\lawyer_consult')
path = out / '杭州幼儿在园受伤律师咨询材料_2026-10-20.docx'
backup = qa / '咨询材料修改前六页版.docx'
if not backup.exists():
    shutil.copy2(path, backup)
doc = Document(path)
for element in list(doc._element.body):
    if element.tag != qn('w:sectPr'):
        doc._element.body.remove(element)

normal = doc.styles['Normal']
normal.font.size = Pt(11.5)
normal.paragraph_format.line_spacing = 1.14
normal.paragraph_format.space_after = Pt(6)
lang = OxmlElement('w:lang')
lang.set(qn('w:val'), 'zh-CN')
lang.set(qn('w:eastAsia'), 'zh-CN')
normal._element.get_or_add_rPr().append(lang)
kinsoku = OxmlElement('w:kinsoku')
kinsoku.set(qn('w:val'), '1')
normal._element.get_or_add_pPr().append(kinsoku)
doc.styles['Title'].font.size = Pt(18)
doc.styles['Title'].paragraph_format.space_after = Pt(6)
doc.styles['Heading 1'].font.size = Pt(13)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(8)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(6)
for style in doc.styles:
    for border in list(style._element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
for name in ['Title', 'Heading 1', 'Normal']:
    doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)

plain = []
def para(text, style=None, bold=False):
    p = doc.add_paragraph(style=style)
    p.add_run(text).bold = bold
    p.paragraph_format.widow_control = True
    plain.append(text)
    return p

para('杭州幼儿在园受伤线上律师咨询', 'Title')
p = para('情况截至2026年10月20日    幼儿园全称及园区：【待补充】')
p.runs[0].font.size = Pt(9)

para('一  事件经过', 'Heading 1')
para('孩子男，事发时约2岁4个月。2026年10月15日约11时15分，在杭州幼儿园托二班教室额头受伤，伤口约2.5厘米。约11时30分老师来电，最初告知孩子摔倒；之后我们到园查看监控，认为另一同园幼儿的动作直接导致受伤。具体致伤点、事发当刻老师人数及是否及时制止，仍需结合完整监控核实。')
para('当天在浙江大学医学院附属儿童医院清创缝合，家长记录缝7针；10月18日又到杭州整形医院拆线后重新全麻美容缝合，未住院。首诊病历已提示永久性瘢痕、影响外观及后期整形等风险。')
para('10月19日，双方家长、园方、教育指导中心人员及民警参与沟通。家长记录民警已保存监控，备份范围及副本获取方式尚需确认。我们已有病历、伤口照片、部分票据、费用台账和3段原始录音，可按需要提供。')

para('二  目前现状和总费用', 'Heading 1')
para('目前已发生医疗费约1万元，包括两次缝合及相关检查、麻醉、药品等，准确金额待票据汇总。不含后续激光、祛疤药物及其他瘢痕修复费用；后续治疗次数、周期和总费用尚未确定。护理、交通、营养等费用尚未统一核算。', bold=True)
para('当前主要担心额头留疤和后续费用承担。现有资料未见明确的书面责任比例、最终赔偿金额及分阶段付款安排。已收到的垫付款、已签文件或收据情况：【待补充】。')

para('三  录音总结', 'Heading 1')
para('10月16日园长谈话：园方表示积极处理、先治疗并保留票据，谈到现阶段费用垫付；后续治疗听医生意见。家长希望将长期跟进及费用安排落实为机构层面的书面文件，尚未形成明确协议。')
para('10月19日园长谈话：园方表示不推脱应承担的责任，在责任范围内会支付费用；以公办园流程及保险处理为由，主张治疗结束后一次性结算。家长要求分阶段处理费用并提供书面事故说明，尚未落实具体方案。')
para('10月19日韩老师等人谈话：录音有针对当时现有两次缝合费用的保险不足部分承担表态；韩老师表示可以继续协调，并提出经济困难可申请垫付。具体责任比例、后续全部费用及付款时间未确定。')
p = para('录音说明：以上依现有机器转写概括，关键表态的说话人、权限和范围需核听原音。“三万元”等为举例，尚未议定赔偿总额。')
p.runs[0].font.size = Pt(9)

doc.add_page_break()
para('四  希望咨询的问题', 'Heading 1')
questions = [
('1. 责任如何划分？', '幼儿园和另一幼儿监护人分别承担什么责任，应向谁主张，是否需要一起起诉？目前最需要补强哪些证据？'),
('2. 能否先处理已发生费用？', '必须等全部治疗结束才能协商、赔偿或走保险吗？能否先支付现有约1万元医疗费，并分阶段结算后续合理费用？'),
('3. 怎样保障后续祛疤费用？', '激光、祛疤药物、瘢痕修复及必要麻醉费用，需什么医疗意见或鉴定？协议怎样写，才能保留暂未确定的后续费用，避免被一次性结清？'),
('4. 二次美容缝合费用能否主张？', '第一次普通缝合后，再拆线做全麻美容缝合，怎样证明合理性和必要性？如果对方以重复治疗或保险不赔为由拒付，应怎样处理？'),
('5. 护理费和工作损失怎样算？', '父亲请带薪假，工资未必减少；母亲经营公司。合理护理期、人数、日标准怎样确定，母亲因照护减少的经营收入怎样举证，哪些项目不能重复计算？'),
('6. 其他赔偿可以提出多少？', '营养费是否需要医嘱或购物发票？交通费、精神损害抚慰金等在杭州有什么支持条件？请结合孩子年龄、面部伤口和两次治疗，区分协商报价与法院较可能支持的范围。'),
('7. 是否需要鉴定？', '额头约2.5厘米伤口、未住院的情况，是否需要或值得做伤残、护理期及营养期等鉴定？何时进行较合适，未构成伤残是否影响精神损害主张？'),
('8. 录音和下一步处理怎样安排？', '现有费用承担等口头表态能证明什么？现阶段应先协商、调解还是诉讼？请指出最应先做的事项，以及初步咨询后需要补交的关键材料。')
]
for title, text in questions:
    p = doc.add_paragraph()
    p.add_run(title).bold = True
    p.add_run(text)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.widow_control = True
    plain.append(title + text)

p = para('希望得到的意见：责任判断、赔偿项目及大致金额范围、后续费用约定方式、优先补证及处理步骤。另请说明后续委托范围和收费。')
p.paragraph_format.space_before = Pt(6)

doc.save(path)
textpath = out / '首次联系杭州律师_微信简版.txt'
textpath.write_text('\n\n'.join(plain) + '\n', encoding='utf-8-sig')
print(json.dumps({'docx':str(path),'text':str(textpath),'chars':sum(map(len,plain))},ensure_ascii=False))
