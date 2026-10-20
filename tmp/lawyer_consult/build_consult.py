from pathlib import Path
from decimal import Decimal
import hashlib
import json
import re

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(r'C:\dev\XuChe\幼儿受伤处理记录')
OUT = ROOT / '08-律师咨询'
QA = Path(r'C:\dev\XuChe\tmp\lawyer_consult')
OUT.mkdir(parents=True, exist_ok=True)
QA.mkdir(parents=True, exist_ok=True)
source_files = sorted(p for p in ROOT.rglob('*') if p.is_file() and '08-律师咨询' not in p.parts)
manifest = [{'path': str(p.relative_to(ROOT)), 'bytes': p.stat().st_size,
             'sha256': hashlib.sha256(p.read_bytes()).hexdigest()} for p in source_files]
(QA / 'source_manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(.65)
sec.bottom_margin = Inches(.65)
sec.left_margin = Inches(.75)
sec.right_margin = Inches(.75)
sec.header_distance = Inches(.25)
sec.footer_distance = Inches(.25)

for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Header', 'Footer']:
    st = doc.styles[name]
    st.font.name = 'Microsoft YaHei' if name != 'Normal' else 'SimSun'
    st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '微软雅黑' if name != 'Normal' else '宋体')
    st.font.color.rgb = RGBColor(0, 0, 0)
for st in doc.styles:
    for border in list(st._element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing = 1.16
doc.styles['Normal'].paragraph_format.space_after = Pt(5)
doc.styles['Title'].font.size = Pt(20)
doc.styles['Title'].paragraph_format.space_after = Pt(8)
doc.styles['Heading 1'].font.size = Pt(15)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(0)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(9)
doc.styles['Heading 2'].font.size = Pt(11)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(6)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(4)

header = sec.header.paragraphs[0]
header.text = '杭州幼儿在园受伤律师咨询材料'
header.runs[0].font.size = Pt(8)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = footer.add_run('2026年10月20日整理  第 ')
run.font.size = Pt(8)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
footer.add_run(' 页').font.size = Pt(8)

plain_sections = []
def p(text, bold=False, small=False):
    para = doc.add_paragraph()
    para.paragraph_format.widow_control = True
    r = para.add_run(text)
    r.bold = bold
    if small:
        r.font.size = Pt(9)
    plain_sections.append(text)
    return para

def heading(text, level=1):
    para = doc.add_paragraph(text, style=f'Heading {level}')
    plain_sections.append('\n' + text)
    return para

def page(text):
    doc.add_page_break()
    heading(text)

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for cell, label, width in zip(t.rows[0].cells, headers, widths):
        cell.text = label
        cell.width = Inches(width)
    repeat = OxmlElement('w:tblHeader')
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells = t.add_row().cells
        for cell, text, width in zip(cells, row, widths):
            cell.text = str(text)
            cell.width = Inches(width)
    for idx, row in enumerate(t.rows):
        cant = OxmlElement('w:cantSplit')
        row._tr.get_or_add_trPr().append(cant)
        for col, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for edge in ['top', 'left', 'bottom', 'right']:
                e = OxmlElement('w:' + edge)
                e.set(qn('w:val'), 'single')
                e.set(qn('w:sz'), '4')
                e.set(qn('w:color'), 'D9D9D9')
                borders.append(e)
            pr.append(borders)
            margins = OxmlElement('w:tcMar')
            for edge in ['top', 'left', 'bottom', 'right']:
                e = OxmlElement('w:' + edge)
                e.set(qn('w:w'), '75')
                e.set(qn('w:type'), 'dxa')
                margins.append(e)
            pr.append(margins)
            if idx == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'E7EDF3')
                pr.append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(1)
                para.paragraph_format.line_spacing = 1.08
                if col == 0:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in para.runs:
                    r.font.name = 'SimSun'
                    r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')
                    r.font.size = Pt(10)
                    r.bold = idx == 0
                    r.font.color.rgb = RGBColor(0, 0, 0)
    plain_sections.append(' | '.join(headers))
    plain_sections.extend(' | '.join(map(str, row)) for row in rows)
    return t

def link(label, url):
    para = doc.add_paragraph()
    rel = para.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'), rel)
    r = OxmlElement('w:r')
    pr = OxmlElement('w:rPr')
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '18'); pr.append(sz)
    color = OxmlElement('w:color'); color.set(qn('w:val'), '000000'); pr.append(color)
    r.append(pr)
    txt = OxmlElement('w:t'); txt.text = label; r.append(txt)
    h.append(r); para._p.append(h)
    para.paragraph_format.space_after = Pt(2)
    plain_sections.append(label + '：' + url)

# PAGE 1
doc.add_paragraph('杭州幼儿在园受伤律师咨询材料', style='Title')
plain_sections.append('杭州幼儿在园受伤律师咨询材料')
p('整理日期  2026年10月20日', small=True)
p('我们希望律师根据现有证据，评估幼儿园及另一幼儿监护人的责任、已发生与后续治疗费用的主张方式，并给出杭州地区的赔偿测算和下一步方案。现阶段尤其希望明确阶段性付款、后续祛疤费用保留及需要立即补强的证据。')
heading('一  基本情况', 2)
p('孩子：徐澈，男。2026年10月15日儿保病历记载年龄为2岁4个月；准确出生日期及监护关系证明待补。')
p('事发：家长台账记载2026年10月15日约11时15分，杭州某幼儿园托二班教室；约11时30分接到老师电话。幼儿园全称、园区、具体地址及运营主体待补。园方谈话中自称公办园，尚缺登记或招生收费材料核实。')
p('伤情：家长说明额头伤口约2.5厘米；照片可见伤口接近眉间。首诊诊断为皮肤裂伤。家长说明未住院；台账记载先缝合7针，10月18日再次全麻美容缝合。后一次手术病历、麻醉记录和主要费用票据尚未见于当前资料。')
p('事因：家长查看监控后认为另一同园幼儿的动作直接导致受伤。当前工作空间没有监控视频，具体动作、哪次倒下致伤、接触物及老师是否及时制止，仍须核实。本文不将“故意欺负”或“园方全责”列为已查明结论。')
heading('二  本次希望律师先回答的事项', 2)
p('1. 应向哪些主体主张、各自承担何种责任，是否适合共同起诉；监控与人员安排需要补哪些证据。')
p('2. 已花医疗费家长称约1万元，后续激光费用未确定。能否现在先处理已发生费用，并明确保留后续必要治疗费用。')
p('3. 请分别给出协商报价、法院可能支持的金额范围及各项证据要求，不将此前聊天中的估算当作既定赔偿金额。')
heading('三  事实来源说明', 2)
p('病历、票据可核对就诊及金额；台账反映家长记录；谈话转写仅作为检索录音的线索。文中标注“待补”“待核”的事项，表示当前材料不足，不表示现实中一定没有该材料或该事实。', small=True)

# PAGE 2
page('事件经过和就诊时间线')
p('以下按病历及票据、家长台账、谈话记录分别注明来源。具体时点发生冲突时保留差异，交由律师结合原始资料核实。', small=True)
timeline = [
('10月15日约11时15分及11时30分', '台账记事发约11时15分，11时30分接到老师电话，最初被告知孩子在教室摔倒。台账还记“事发时只有1名老师、平时配4名”，但谈话中另有厕所照护、轮流吃饭及2人、3人的不同说法，人数不能据台账直接定案。'),
('10月15日13时06分及14时25分', '儿保急诊、烧伤瘢痕整形门诊病历有就诊时间。诊断皮肤裂伤；急诊记录局部麻醉、清创缝合，并告知永久性瘢痕影响外观、后期整形等风险。整形门诊记伤口已缝合包扎，建议每天换药、1周复查。台账写首次缝7针及12时就医，针数和流程时点需结合手术记录再核。'),
('10月15日下午及晚间', '台账记多位老师陪同就医，约18时托二班4位老师上门探望；具体人员、身份及各段陪同时间待核。'),
('10月16日', '台账记园长上午探望、家长12时40分到园看监控且不允许拍摄，之后知道涉及另一同园幼儿；家长记录当时口头谈到后续费用。社区处方及换药票据可核实当日换药，结算时间为14时53分左右。'),
('10月17日', '家长记录主动联系另一幼儿家长，对方表示此前不清楚；此点应补聊天原始记录。台账记杭州整形医院门诊、换药并提出建群；14元换药票据可核实，25.30元挂号、225.76元抽血暂仅见台账。'),
('10月18日', '台账记杭州整形医院再次拆线、全麻美容缝合，杨老师、虞老师陪同。家长口述手术约八九千元、累计医疗费约1万元。当前没有该次完整病历、麻醉及结算材料，需优先补齐。'),
('10月19日', '台账记约14时双方家长到园、教育指导中心人员及民警参与沟通，民警保存监控。两段面谈文件日期经已有核对报告修订为10月19日。现场有费用、责任及结算方式表态，未见签署的责任或赔偿协议。'),
('截至10月20日当前归档状态', '已有3段录音及机器稿、病历照片、伤口照片、6张医疗电子票据和1张停车发票。未见监控副本、园方盖章事故说明、保单、书面赔偿协议、垫付款到账凭证或正式护理及营养期意见。')
]
for title, text in timeline:
    heading(title, 2); p(text)
p('出处：工作簿“事件台账”“伤情跟踪”；02-医疗记录中的急诊、整形门诊、社区处方及结算照片；05-票据原件；04目录当前合并后的谈话记录。首次病历的“摔伤”系就诊病史，是否与另一幼儿有关仍应结合监控。', small=True)

# PAGE 3
page('希望律师逐项评估的问题')
questions = [
('1  责任主体和适用规则', '孩子未满3岁、就读托二班，应先核实该园及托班的法律主体和管理性质吗？同园幼儿致伤时，园方和另一幼儿监护人分别如何担责、举证，是否应共同列为被告？'),
('2  监控和现场管理', '如何取得完整连续监控、相关机位及原始文件？民警保存范围如何确认？事发前后老师的位置、交接、轮流吃饭、观察和制止是否充分，需要哪些排班、现场及处置记录？'),
('3  先赔已发生费用', '园方提出治疗结束再整体走保险，家长能否先协商或起诉已发生费用？“经济困难报告”是否只是垫付途径？先付款时应怎样写明款项性质及未结清范围？'),
('4  后续祛疤和治疗费用', '必要激光、祛疤药物及材料、相关麻醉和复诊费用，如何证明与此次受伤的关联及必要性？可否凭医疗意见估算，或按阶段处理？协议如何约定新增必要费用、期限及争议解决？'),
('5  二次美容缝合费用', '第一次选择普通缝合后再全麻美容缝合，怎样证明合理性和必要性？需补哪些病历、知情同意书、医生意见及收费明细？对方若主张重复治疗或保险不赔，如何应对？'),
('6  护理费和父母工作损失', '父亲使用带薪假、工资未必减少；母亲经营公司，登记形态和收入待补。杭州怎样认定额外护理、护理期及人数？经营收入损失需哪些证据，怎样避免与护理费重复？'),
('7  营养费及就医交通', '未住院情况下，营养必要性和期限如何证明，是否需购物发票？同城交通如何计算？现有25元小区停车费与就医关联不直接，是否可主张、需要怎样补证？'),
('8  精神损害和伤残评估', '未评伤残的幼儿面部留疤，在杭州的支持条件和金额范围如何？何时、是否值得鉴定？0—6周岁面部瘢痕长度面积按50%计算的规则，怎样结合最终位置、宽度和面积适用？'),
('9  缺课退费及返园安排', '因伤未入园的伙食费、保育费分别如何退，退费与侵权损失如何衔接？需要哪些收费文件和缺课医嘱？返园照护、伤口保护及家园沟通能否形成书面安排？'),
('10  录音和口头表态效力', '现有费用保险不足部分承担等表态，如何确认说话人、代表权限及承诺范围？既有谈话中家长说“可以”“其他费用后谈”，是否影响尚未确定项目的主张？需要怎样制作证据说明？'),
('11  程序及费用收益', '应先律师函、人民调解还是诉讼？由哪个法院管辖？诉讼时效从何时计算及如何保留证据？请区分咨询、文书、协商、诉讼、鉴定费用及可能获赔金额。'),
('12  明确此次咨询的产出', '请列出各项赔偿金额或计算公式、待补证据、协商报价和可能裁判区间；指出当前最应做的3件事，并评估拟签协议，尤其是后续费用是否被一并结清。')
]
for title, text in questions:
    para = p(title + '。' + text)
    para.clear()
    para.add_run(title + '。').bold = True
    para.add_run(text)

# PAGE 4
page('费用核对和待补金额')
p('家长目前说明医疗费累计约1万元。该数额尚未由当前文件全部核验，主要缺10月18日手术及麻醉票据。下表按现有原票核对，未将同一支出的照片、结算单和电子票据重复相加。')
fee_rows = [
('10月15日', '儿保急诊挂号', '35.00', '05-001电子票据，金额一致'),
('10月15日', '清创缝合及药品', '466.10', '05-002电子票据；台账写466.00，少0.10'),
('10月15日', '儿保整形门诊挂号', '80.00', '05-003原票及导诊单均为10月15日；PDF文件名写10月16日'),
('10月16日', '社区挂号', '5.00', '05-004电子票据'),
('10月16日', '社区换药及药品', '16.45', '05-005电子票据；照片中15元换药及1.45元药费属于同一笔'),
('10月17日', '杭州整形医院换药', '14.00', '05-008电子发票；02目录结算单为同一笔'),
('小计', '6张医疗电子票据', '616.55', '这是票据总额小计，不等于全部自费或最终获赔额'),
('10月17日', '挂号及抽血', '251.06', '台账25.30＋225.76；当前未见对应票据或检查材料'),
('10月17日付款', '小区停车', '25.00', '付款截图＋同项发票；10月18日开票，记租赁期间10月16—17日，是否属于就医交通待证'),
('合计', '已录台账按原票修正后', '892.61', '医疗867.61＋停车25.00；仍不代表全部支出')
]
table(['日期', '费用项目', '金额 元', '核对结果'], fee_rows, [.85, 1.6, .85, 3.7])
heading('应补齐的费用及付款情况', 2)
p('优先补10月18日手术、麻醉、材料、药品、术前检查的发票及结算清单，再补后续换药、复诊、祛疤药物和交通。每笔标明付款人、付款方式、退款或报销、园方或老师是否已垫付及凭证。')
p('部分医疗票据显示个人账户或家庭共济、其他支付或补助。应请律师区分费用总额、实际支付和补偿来源，不能把票据总额直接写成“全部现金自费”。')
p('护理期、营养期、后续激光次数及费用尚无完整医疗意见；这些项目暂留待律师测算。无住院记录，当前不把住院伙食补助列成已发生损失。')
p('出处：工作簿“费用明细”D2:D13；05目录6份PDF及停车发票；02目录就诊、结算和付款截图。仅作核对，不修改原台账或原文件。', small=True)

# PAGE 5
page('录音中需要律师核听的重点')
p('共有3段原始M4A。两段10月19日录音合计约127分钟；另有“10-16园长”约22分钟，当前整理记录标注为10月16日，举证时仍请家长确认录制时间。3份配套TXT已合并要点及转写正文。本次读取文本和已有核对信息，并核验原音频校验值，未对约149分钟原音频逐字人工复听。')
p('下列是检索线索，不是已确认的逐字引文或责任结论。请律师结合原音确认完整上下文、说话人及其身份，再判断证明力。', bold=True)
audio_rows = [
('10-16园长.m4a', '02:10—07:51\n16:00—19:20', '治疗和费用垫付、听医生意见、出具机构证明及持续跟进。当前配套记录标注谈话日期为10月16日。'),
('园长(1).m4a', '03:34—05:07\n59:43—60:39', '现场人员、厕所照护、轮流吃饭及抬床；区分总配置、镜头可见、事发当刻和事后返回人数。'),
('园长(1).m4a', '19:00—21:13', '具体致伤点尚不清楚；讨论衣物排查及书面事故说明。未见说明已交付或确定交付日。'),
('园长(1).m4a', '24:07—24:17\n53:41—54:22', '责任范围内费用承担，以及园方称结束后一次性结算、家长提出分阶段处理。未确定付款日期。'),
('韩领导(1).m4a', '11:50—14:14', '当时现有两次缝合费用、保险不能报部分的承担表态、垫付报告及继续协调；需核说话人和代表权限。'),
('韩领导(1).m4a', '18:19—19:28\n22:01—22:38', '调解中的责任意见及假设花费“三万元”。不是已议定赔偿3万元，也不是确定比例或无限兜底。'),
('韩领导(1).m4a', '28:32—30:51\n38:18—41:07', '书面材料、整体调解、结案后新增费用；另涉及通知对方家长、人员安排及继续排查。')
]
table(['原音文件', '相对时间', '核听事项'], audio_rows, [1.5, 1.35, 4.15])
heading('转写的已知局限', 2)
p('当前合并记录已提示：机器稿存在错词、方言、重叠说话、否定语义疑误及说话人不明。当前要点纠正了旧索引扩大承诺、错位时间和将口头意见写成正式结论的问题。引用时应结合各文件前部的核对要点及原音。')
p('不把“这是我管理的问题”直接写成承认全部致伤责任；不把另一家长附条件发言写成承认打人；不把误识别的严重措辞写成威胁；不把他人治疗案例及估计写成本孩子医嘱。', small=True)
p('当前文字文件：04-001_2026-10-19_园长面谈录音转写.txt；04-003_2026-10-19_韩领导谈话录音转写.txt；04-007_10-16_园长谈话录音转写.txt。均位于04-沟通记录目录，各自前部为要点和疑点提示。', small=True)

# PAGE 6
page('咨询前需要补充和携带的材料')
heading('优先补充', 2)
checklist = [
'1. 幼儿园完整名称、园区及地址、运营或登记主体、入园合同和缴费凭证；确认托二班的管理性质及对应责任主体。',
'2. 孩子准确出生日期、父母姓名和联系方式、监护关系证明。正式委托或办案所需身份证明按律师要求提供。',
'3. 完整监控：机位、连续覆盖时段、能否导出原文件；谁保存、民警备份范围和获取方式；保存申请及回应的微信记录。',
'4. 事故及处置原始记录：发现受伤、通知、送医时间；老师位置和分工；园方事故说明及现场照片。明确哪些是亲见、哪些是听说。',
'5. 两次手术完整材料，尤其10月18日病历、手术及麻醉记录、同意书、发票和收费清单；补首次缝针数和伤口长度的医疗记录。',
'6. 后续治疗意见：瘢痕风险、祛疤项目、必要性、预计次数和费用、护理及营养需要；这些应由医生据实出具。',
'7. 全部支付、垫付及报销记录；父亲请假和工资情况，母亲营业执照、实际经营及收入变化材料；实际额外护理天数。',
'8. 对方幼儿及监护人的身份、联系信息、孩子年龄；10月17日联系记录及双方原话，避免只留家长二次概括。',
'9. 园方、教育指导中心和民警的姓名、单位及处理记录；保单、保险公司、险种、报案号和理赔条件；已签文件或收款收据。',
'10. 因伤缺课天数、收费及退费规则、返园医嘱与保护需求；说明此前是否有类似情况及园方是否知情，若无则不添加。'
]
for item in checklist:
    p(item)
heading('第一次联系律师时怎么发', 2)
p('先发送同目录的微信简版，告知希望咨询的范围及有无委托需求。面谈时提供本材料、当前工作簿、关键病历与费用清单；录音提供原文件及对应索引。请律师列明待补材料、书面意见内容、收费和下一步时点。')
heading('本次资料读取范围', 2)
p('当前原资料共30份：1个工作簿、6份PDF、10张图片、3段原音频及10份TXT。已通读文字内容并核对最新合并稿；PDF逐份提取核对金额，图片逐张查看。原音频已读取计算校验值，内容以现有转写及核对信息为线索，未作全程人工听辨。', small=True)
p('当前07-园方材料、06-证人名单目录仅有提示文本，01-事发经过只有未填写模板。当前未归档的材料请补交律师；不据此认定有关证据不存在。本次仅新建咨询材料，不改动原资料。', small=True)
heading('供律师核对的法规入口', 2)
link('民法典 责任主体及精神损害等条款', 'https://www.cac.gov.cn/2020-06/01/c_15925617772683196.htm')
link('最高法 人身损害赔偿司法解释 第6至12条', 'https://www.court.gov.cn/zixun/xiangqing/357071.html')
link('人体损伤致残程度分级 第5.10.2及6.11条', 'https://sfj.xlgl.gov.cn/sfj/ywgz/ggflfw/sfjd/2cad4f41d39d4fe6a232ac8c3a670277/index.html')

main_doc = OUT / '杭州幼儿在园受伤律师咨询材料_2026-10-20.docx'
doc.save(main_doc)

wechat = '''律师您好，想咨询杭州幼儿在园受伤的责任和赔偿问题。

孩子男，事发时约2岁4个月。2026年10月15日在幼儿园托二班教室额头受伤，家长记录伤口约2.5厘米。最初老师电话说摔倒，后来我们查看监控后认为与另一同园幼儿的动作直接有关，具体致伤点及责任尚未查清。

当天到浙江大学医学院附属儿童医院清创缝合，家长记录缝7针；10月18日又到杭州整形医院全麻美容缝合，未住院。首诊病历告知永久性瘢痕、影响外观及后期整形等风险。医疗费目前约1万元，手术票据正在补齐，后续祛疤方案和总费用未确定。

10月19日双方家长、园方及教育指导中心人员、民警参与沟通。园方谈到责任范围内费用及当时现有费用的保险不足部分承担，但尚无明确书面赔付方案，提出等治疗结束再整体处理。家长台账记民警已保存监控，我们尚需确认备份范围和取得副本。现场老师人数有不同说法，需要结合完整监控核实。

已有事件和费用台账、病历及伤口照片、部分票据、3段原始录音和机器转写。转写并非逐字核准稿，可配合原音核听。

重点想请您评估：
1. 幼儿园与另一幼儿监护人分别怎样担责，应向谁主张或起诉？
2. 能否先处理已发生费用，怎样书面保留后续必要祛疤、激光及相关费用？
3. 二次美容缝合、全麻费用的合理性需要哪些证据？
4. 护理费如何计算？父亲请带薪假，母亲经营公司，照护造成的收入损失怎样举证？
5. 营养费、交通费、精神损害抚慰金在杭州的支持条件和金额范围如何？是否需要、何时考虑瘢痕相关鉴定？
6. 现有录音表态能证明什么，还缺哪些关键证据？协商、调解或诉讼怎样安排最合适？

希望您在看过材料后，区分协商报价与可能裁判金额，并告知咨询及后续服务收费。我整理了详细时间线、费用核对和录音索引，可供进一步查看。

补充后再发送：幼儿园全称及园区【待填】；我的姓名和联系方式【待填】；目前是否收到垫付款、是否签过文件【待填】。
'''
(OUT / '首次联系杭州律师_微信简版.txt').write_text(wechat, encoding='utf-8-sig')
(QA / 'full_text_for_qa.txt').write_text('\n'.join(plain_sections), encoding='utf-8')
print(json.dumps({'docx':str(main_doc),'wechat':str(OUT/'首次联系杭州律师_微信简版.txt'),'source_count':len(source_files),'paragraphs':len(doc.paragraphs)},ensure_ascii=False))
